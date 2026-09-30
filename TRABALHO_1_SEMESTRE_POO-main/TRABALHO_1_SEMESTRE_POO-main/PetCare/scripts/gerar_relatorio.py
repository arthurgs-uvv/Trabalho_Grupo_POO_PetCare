"""Gera o relatorio editavel de duas paginas. Requer python-docx apenas para regenerar."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

raiz = Path(__file__).resolve().parents[1]
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.8)
for nome in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    estilo = doc.styles[nome]
    estilo.font.name = 'Times New Roman'
    estilo.font.color.rgb = RGBColor(0, 0, 0)
normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(17)
doc.styles['Title'].paragraph_format.space_after = Pt(5)
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(4)

def par(texto):
    p = doc.add_paragraph(texto)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def titulo(texto): doc.add_heading(texto, level=1)

doc.add_paragraph('Aplicação de POO no Sistema PetCare', style='Title')
p=doc.add_paragraph('Programação Orientada a Objetos I | Relatório do trabalho em grupo')
p.paragraph_format.space_after=Pt(8)
p.runs[0].font.size=Pt(10)
for i in range(1,6):
    p=doc.add_paragraph(f'{i}. Nome: __________________________________   Matrícula: ______________')
    p.paragraph_format.space_after=Pt(2)
    p.runs[0].font.size=Pt(10)
p=doc.add_paragraph('Professor(a): __________________________   Data: ____ / ____ / ________')
p.runs[0].font.size=Pt(10)
p.paragraph_format.space_after=Pt(8)
titulo('Organização do sistema')
par('O Sistema Integrado PetCare representa os cadastros, o histórico clínico, os agendamentos, o estoque e o financeiro de uma clínica veterinária. Na implementação, organizamos as classes no pacote petcare.model e usamos petcare.Main para demonstrar a integração. Aplicamos os quatro conceitos solicitados no trabalho: herança, encapsulamento, composição e delegação. As regras são verificadas nos próprios objetos, evitando que uma operação inválida deixe os dados inconsistentes.')
titulo('Herança')
par('A classe Usuario concentra os dados e comportamentos comuns, como nome, e-mail, senha, login e desativação. Administrador, Veterinario, Recepcionista e Tutor utilizam extends Usuario e chamam super no construtor. Cada perfil acrescenta suas responsabilidades: o veterinário registra laudos, a recepcionista gerencia cadastros e agenda, e o administrador ajusta o estoque e consulta a auditoria.')
par('No módulo clínico, Consulta, Cirurgia, Exame, Vacina e Tratamento herdam de RegistroClinico. Dessa forma, reutilizam identificação, data, veterinário, laudo e bloqueio após a finalização. Os métodos exibir são sobrescritos com @Override para mostrar informações específicas. Em Main, listas de Usuario e RegistroClinico recebem objetos das subclasses e chamam exibir. Isso também demonstra polimorfismo: a mesma chamada executa o comportamento do objeto concreto.')
titulo('Encapsulamento')
par('Todos os atributos são private. Os métodos controlam a leitura e a alteração dos dados: Animal.setPeso rejeita valores inválidos; Usuario.setEmail valida o endereço; e Cirurgia.setDuracao impede duração inferior a 120 minutos. ItemEstoque só permite ativar o controle de medicamentos quando lote, validade e responsável estão preenchidos. A validação ocorre antes de alterar o atributo, preservando o valor anterior se houver erro.')
par('As coleções são devolvidas sem permissão de alteração externa. Assim, não é possível apagar consultas diretamente pela lista retornada por getConsultas. HistoricoClinico.finalizar valida os registros e bloqueia novas inclusões; os registros também são finalizados, impedindo mudanças posteriores nos laudos e demais dados clínicos.')

doc.add_page_break()
titulo('Composição')
par('Animal cria seu próprio HistoricoClinico no construtor e mantém essa referência como final. Esse é o exemplo principal de composição: o histórico faz parte do animal e não pode ser substituído por outro objeto. O histórico reúne listas de consultas, vacinas, cirurgias, exames e tratamentos. A inclusão também confere que um registro não seja vinculado a dois animais.')
par('Estoque mantém uma coleção de ItemEstoque; Agendamento possui referências ao animal, ao veterinário e ao notificador; e Fatura referencia o tutor e o notificador. Essas relações representam o uso de objetos como partes ou colaboradores de outros objetos. Em UML estrito, referências compartilhadas a tutor, veterinário e notificador são associações, pois esses objetos existem independentemente. Cada animal mantém apenas um tutor principal, e a inclusão na lista de outro tutor é rejeitada.')
titulo('Delegação')
par('Agendamento coordena a marcação, o cancelamento e o reagendamento, mas delega as comunicações ao Notificador. Quando um agendamento é aceito, chama enviarConfirmacao; quando é cancelado, chama enviarCancelamento. Estoque delega os alertas de quantidade baixa, e Fatura delega o envio das informações de cobrança. O canal EMAIL, SMS ou APP fica concentrado em Notificador, sem repetir essa lógica nos outros módulos.')
par('Também utilizamos delegação entre a agenda e o estoque: o agendamento solicita a reserva dos insumos, e Estoque coordena os itens envolvidos. ItemEstoque controla o saldo físico e a quantidade reservada. Ao cancelar, a reserva é liberada; ao concluir, os materiais reservados são consumidos. Antes de reservar, todos os itens são conferidos para evitar uma reserva parcial.')
titulo('Integração e verificação')
par('A demonstração em Main instancia todas as classes de domínio e executa cadastro, busca, procedimentos clínicos, agenda, estoque, faturas, relatório e auditoria. A agenda verifica a disponibilidade do veterinário e conflitos de intervalo entre veterinário, sala e animal. Todo procedimento deve terminar até as 18h. O reagendamento repete essas verificações e mantém o horário anterior quando a alteração é recusada.')
par('Os testes automatizados conferem as regras obrigatórias, as operações inválidas e os resultados dos relatórios. Na demonstração, quatro atendimentos concluídos geram uma taxa de retorno de 25% e duração média planejada de 52,5 minutos. Duas faturas totalizam R$ 400,00 emitidos e R$ 150,00 recebidos. LogAuditoria registra data, hora, usuário e ação; o estoque mantém o histórico das movimentações por lote e responsável.')
titulo('Limites da versão e referências')
par('Esta versão acadêmica funciona no console e mantém dados em memória durante a execução. Notificações, boleto e pagamento são simulados; confirmarPagOnline retorna true conforme o enunciado. O produto web, o banco de dados e as integrações externas pertencem à evolução do sistema. A entrega aplica os conceitos de POO aos módulos e às regras previstos para a atividade.')
p=par('Referências: Trabalho em Grupo - Sistema Integrado PetCare, seções 3 a 6; Sistema Integrado PetCare - Requisitos e Escopo. Materiais fornecidos na atividade.')
p.runs[0].font.size=Pt(10)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('PetCare | ')
f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
for r in footer.runs:r.font.size=Pt(9)
doc.core_properties.title='Aplicação de POO no Sistema PetCare'
doc.core_properties.author='Grupo PetCare'
doc.core_properties.subject='Herança, encapsulamento, composição e delegação'
# Remove bordas e fontes de tema herdadas do modelo padrao do Word.
for elemento in doc.styles.element.iter():
    if elemento.tag == qn('w:pBdr'):
        elemento.getparent().remove(elemento)
for nome in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
    fontes = doc.styles[nome].element.get_or_add_rPr().get_or_add_rFonts()
    for atributo in list(fontes.attrib):
        if 'Theme' in atributo or 'theme' in atributo:
            del fontes.attrib[atributo]
    fontes.set(qn('w:ascii'), 'Times New Roman')
    fontes.set(qn('w:hAnsi'), 'Times New Roman')
doc.save(raiz/'docs/Relatorio_POO_PetCare.docx')
